import io
from datetime import datetime

from openpyxl import Workbook

from app.services import importer


def make_xlsx(rows, sheet_title="Movimientos") -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_title
    for r in rows:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


STATEMENT = [
    ["Extracto enero", None, None],
    ["Fecha", "Descripción", "Valor"],
    [datetime(2026, 1, 5), "Compra Éxito", -50000],
    [datetime(2026, 1, 6), "Nómina", 3000000.5],
    ["07/01/2026", "Texto como fecha", "-12.500"],
    [datetime(2026, 1, 8), "Con decimales raros", -1234.567],
]


def preview(client, auth, account, content, name="extracto.xlsx"):
    return client.post("/api/imports/preview", headers=auth, files={"file": (name, content, "application/octet-stream")},
                       data={"account_id": str(account["id"])})


def test_read_xlsx_keeps_dates_and_formats_numbers():
    rows = importer.read_xlsx(make_xlsx(STATEMENT))
    assert rows[1] == ["Fecha", "Descripción", "Valor"]
    assert rows[2][0] == datetime(2026, 1, 5) and rows[2][2] == "-50000.00"
    assert rows[3][2] == "3000000.50"
    assert rows[5][2] == "-1234.57"  # nunca se lee como 1.234.567


def test_xlsx_preview_end_to_end(client, auth, bank):
    r = preview(client, auth, bank, make_xlsx(STATEMENT))
    assert r.status_code == 200
    body = r.json()
    assert body["header_row"] == 1 and body["stats"]["valid"] == 4
    got = [(x["date"], x["kind"], float(x["amount"])) for x in body["rows"]]
    assert got == [("2026-01-05", "expense", 50000.0), ("2026-01-06", "income", 3000000.5),
                   ("2026-01-07", "expense", 12500.0), ("2026-01-08", "expense", 1234.57)]


def test_xlsx_with_debit_credit_columns_and_empty_first_sheet(client, auth, bank):
    wb = Workbook()
    wb.active.title = "Vacía"
    ws = wb.create_sheet("Datos")
    for row in [["Fecha", "Detalle", "Débito", "Crédito"], [datetime(2026, 2, 1), "Retiro", 100000, None],
                [datetime(2026, 2, 2), "Abono", None, 250000]]:
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    body = preview(client, auth, bank, buf.getvalue()).json()
    assert [(x["kind"], float(x["amount"])) for x in body["rows"]] == [("expense", 100000.0), ("income", 250000.0)]


def test_xlsx_rejects_fake_corrupt_and_legacy_files(client, auth, bank):
    assert preview(client, auth, bank, b"esto es texto, no un excel").status_code == 422
    assert preview(client, auth, bank, b"PK\x03\x04basura-que-no-es-un-zip-valido").status_code == 422
    assert preview(client, auth, bank, b"\xd0\xcf\x11\xe0 old xls", "viejo.xls").status_code == 422


def test_xlsx_row_limit(client, auth, bank):
    rows = [["Fecha", "Descripción", "Valor"]] + [[datetime(2026, 1, 1), f"m{i}", -1000 - i] for i in range(2100)]
    r = preview(client, auth, bank, make_xlsx(rows))
    assert r.status_code == 422 and "más de" in r.json()["detail"]
