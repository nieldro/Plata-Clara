import io
from datetime import date, timedelta
from decimal import Decimal

from openpyxl import load_workbook

from app.services import exporter
from conftest import category_id

TODAY = date.today()


def add(client, auth, account, amount, note="", kind="expense", day=None, category=None, tags=None):
    r = client.post("/api/transactions", headers=auth, json={
        "account_id": account["id"], "category_id": category, "kind": kind, "amount": amount, "note": note,
        "date": (day or TODAY).isoformat(), "tags": tags or []})
    assert r.status_code == 201
    return r.json()


def seed(client, auth, bank):
    mercado = category_id(client, auth, "Mercado")
    add(client, auth, bank, "50000.50", "Éxito", category=mercado, tags=["casa"])
    add(client, auth, bank, 3_000_000, "Nómina", kind="income", day=TODAY - timedelta(days=40))
    add(client, auth, bank, 12_000, "Taxi")


def export(client, auth, fmt="csv", **params):
    return client.get("/api/export/transactions", headers=auth, params={"format": fmt, **params})


def test_csv_opens_in_spanish_excel_and_has_exact_amounts(client, auth, bank):
    seed(client, auth, bank)
    r = export(client, auth)
    assert r.status_code == 200 and r.headers["x-row-count"] == "3"
    assert r.headers["content-type"].startswith("text/csv") and "attachment" in r.headers["content-disposition"]
    raw = r.content
    assert raw.startswith(b"\xef\xbb\xbf")  # BOM: Excel respeta las tildes
    lines = raw.decode("utf-8-sig").splitlines()
    assert lines[0] == "Fecha;Cuenta;Categoría;Tipo;Monto;Nota;Etiquetas;Recibos"
    exito = next(l for l in lines if "Éxito" in l)
    assert ";Gasto;-50000,50;" in exito and exito.endswith(";casa;")           # coma decimal, gasto en negativo
    assert any(";Ingreso;3000000,00;Nómina" in l for l in lines)


def test_formula_injection_is_neutralised_in_csv_and_xlsx(client, auth, bank):
    add(client, auth, bank, 1000, '=HYPERLINK("http://malo.example","clic")')
    add(client, auth, bank, 1000, "+cmd|' /C calc'!A0")
    add(client, auth, bank, 1000, "@SUM(1+1)")
    csv_lines = export(client, auth).content.decode("utf-8-sig")
    assert "'=HYPERLINK" in csv_lines and ";=HYPERLINK" not in csv_lines
    assert ";'+cmd" in csv_lines and ";'@SUM" in csv_lines

    ws = load_workbook(io.BytesIO(export(client, auth, "xlsx").content))["Movimientos"]
    notes = [ws.cell(row=i, column=6) for i in range(2, 5)]
    assert all(c.data_type == "s" and str(c.value).startswith("'") for c in notes)


def test_xlsx_has_real_numbers_dates_and_summary_sheet(client, auth, bank):
    seed(client, auth, bank)
    wb = load_workbook(io.BytesIO(export(client, auth, "xlsx").content))
    assert wb.sheetnames == ["Movimientos", "Resumen"]
    ws = wb["Movimientos"]
    assert [c.value for c in ws[1]] == exporter.HEADERS and ws.freeze_panes == "A2"
    first = [c.value for c in ws[2]]  # el más antiguo: la nómina de hace 40 días
    assert first[3] == "Ingreso" and first[4] == 3_000_000 and first[0].date() == TODAY - timedelta(days=40)
    assert isinstance(ws.cell(row=2, column=5).value, (int, float))  # número, no texto

    res = {r[0].value: r[1].value for r in wb["Resumen"].iter_rows(min_row=2) if r[0].value}
    assert res["Ingresos"] == 3_000_000 and abs(res["Gastos"] - 62_000.5) < 0.01
    assert abs(res["Balance"] - (3_000_000 - 62_000.5)) < 0.01
    assert abs(res["Mercado"] - 50_000.5) < 0.01 and res["Sin categoría"] == 12_000


def test_export_applies_the_same_filters_as_the_history(client, auth, bank):
    seed(client, auth, bank)
    assert export(client, auth, q="taxi").headers["x-row-count"] == "1"
    assert export(client, auth, kind="income").headers["x-row-count"] == "1"
    assert export(client, auth, tag="casa").headers["x-row-count"] == "1"
    assert export(client, auth, date_from=(TODAY - timedelta(days=7)).isoformat()).headers["x-row-count"] == "2"
    assert export(client, auth, min_amount=20000, max_amount=100000).headers["x-row-count"] == "1"
    assert export(client, auth, date_from="2026-12-31", date_to="2026-01-01").status_code == 422


def test_transfers_and_loans_are_labelled_and_left_out_of_the_summary(client, auth, bank):
    other = client.post("/api/accounts", headers=auth, json={"name": "Nequi", "type": "wallet", "institution": "nequi", "initial_balance": 0}).json()
    client.post("/api/transfers", headers=auth, json={"from_account_id": bank["id"], "to_account_id": other["id"], "amount": 7000, "date": TODAY.isoformat()})
    client.post("/api/loans", headers=auth, json={"person": "Ana", "direction": "lent", "amount": 9000, "date": TODAY.isoformat(), "account_id": bank["id"]})
    add(client, auth, bank, 100, "real")
    wb = load_workbook(io.BytesIO(export(client, auth, "xlsx").content))
    labels = sorted(r[3].value for r in wb["Movimientos"].iter_rows(min_row=2))
    assert labels == ["Gasto", "Préstamo", "Transferencia", "Transferencia"]
    res = {r[0].value: r[1].value for r in wb["Resumen"].iter_rows(min_row=2) if r[0].value}
    assert res["Gastos"] == 100 and res["Ingresos"] == 0


def test_exported_csv_can_be_imported_back(client, auth, bank):
    seed(client, auth, bank)
    csv_bytes = export(client, auth).content
    r = client.post("/api/imports/preview", headers=auth, files={"file": ("mis-datos.csv", csv_bytes, "text/csv")},
                    data={"account_id": str(bank["id"])})
    body = r.json()
    assert r.status_code == 200 and body["needs_mapping"] is False and body["stats"]["valid"] == 3
    got = sorted((x["kind"], float(x["amount"])) for x in body["rows"])
    assert got == [("expense", 12000.0), ("expense", 50000.5), ("income", 3000000.0)]


def test_export_requires_login_validates_format_and_is_private(client, auth, bank):
    seed(client, auth, bank)
    assert client.get("/api/export/transactions").status_code == 401
    assert export(client, auth, "pdf").status_code == 422
    client.post("/api/auth/register", json={"email": "o@o.co", "password": "12345678", "name": "Otro"})
    tok = client.post("/api/auth/login", data={"username": "o@o.co", "password": "12345678"}).json()["access_token"]
    other = client.get("/api/export/transactions", headers={"Authorization": f"Bearer {tok}"}, params={"format": "csv"})
    assert other.headers["x-row-count"] == "0" and "Éxito" not in other.content.decode("utf-8-sig")


def test_too_many_rows_asks_to_narrow_filters(client, auth, bank, monkeypatch):
    seed(client, auth, bank)
    monkeypatch.setattr(exporter, "MAX_ROWS", 2)
    r = export(client, auth)
    assert r.status_code == 422 and "Acota" in r.json()["detail"]


def test_safe_text_only_touches_dangerous_starts():
    assert exporter.safe_text("normal") == "normal" and exporter.safe_text("") == "" and exporter.safe_text(None) == ""
    assert exporter.safe_text("=1+1") == "'=1+1" and exporter.safe_text("-5") == "'-5"
    assert exporter.to_csv([exporter.Row(TODAY, "A", "B", "income", "Ingreso", Decimal("1.5"), "x", "", 0)]).count(b"\r\n") == 2
