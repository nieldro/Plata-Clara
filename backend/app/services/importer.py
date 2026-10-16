"""Lectura de extractos bancarios (CSV y, vía `readers`, otros formatos) hacia movimientos.

Pensado para extractos colombianos: delimitador `;` o `,`, tildes, codificación Windows-1252,
montos como «$ 1.234.567,89» o «(50.000)», fechas día/mes/año y columnas débito/crédito separadas.
Nada aquí toca la base de datos: solo convierte texto en filas validadas.
"""
import csv
import datetime as _dt  # alias: un campo llamado «date» con valor por defecto taparía el tipo date
import hashlib
import io
import re
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal, InvalidOperation

from .. import clock
from .rules import fold

MAX_ROWS = 2000
HEADER_SCAN_ROWS = 20

SYNONYMS = {
    "date": ["fecha", "date", "fecha transaccion", "fecha de transaccion", "fecha movimiento",
             "fecha de movimiento", "fecha operacion", "fecha de la operacion", "dia"],
    "debit": ["debito", "debitos", "cargo", "cargos", "retiro", "retiros", "egreso", "egresos", "salida",
              "debit", "valor debito", "valor retiro"],
    "credit": ["credito", "creditos", "abono", "abonos", "deposito", "depositos", "ingreso", "ingresos",
               "entrada", "credit", "valor credito", "valor abono"],
    "amount": ["valor", "monto", "importe", "amount", "valor transaccion", "valor de la transaccion", "total"],
    "description": ["descripcion", "detalle", "concepto", "description", "referencia", "movimiento", "nota", "notas",
                    "transaccion", "observaciones", "comercio", "establecimiento", "nombre"],
}
FIELD_ORDER = ["date", "debit", "credit", "amount", "description"]  # el orden evita que «valor débito» caiga en monto


class ImportError_(ValueError):
    """Error que se le puede mostrar al usuario tal cual."""


# ---------------------------------------------------------------- lectura de tablas
def _decode(data: bytes) -> str:
    for enc in ("utf-8-sig", "cp1252"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("latin-1")


def read_csv(data: bytes) -> list[list[str]]:
    text = _decode(data)
    lines = [ln for ln in text.splitlines() if ln.strip()][:10]
    if not lines:
        raise ImportError_("El archivo está vacío")
    counts = {d: sum(ln.count(d) for ln in lines) for d in (";", ",", "\t", "|")}
    delim = max(counts, key=counts.get) if max(counts.values()) else ","
    reader = csv.reader(io.StringIO(text), delimiter=delim)
    return [[c.strip() for c in row] for row in reader]


def _xlsx_cell(v):
    """Celdas de Excel → algo que el parser entiende. Los números se pasan con 2 decimales para que
    1234.567 no se lea como «1.234.567»; las fechas se conservan como fecha."""
    if v is None:
        return ""
    if isinstance(v, (_dt.datetime, _dt.date)):
        return v
    if isinstance(v, bool):
        return str(v)
    if isinstance(v, (int, float, Decimal)):
        return format(Decimal(str(v)).quantize(Decimal("0.01")), "f")
    return str(v).strip()


def read_xlsx(data: bytes) -> list[list]:
    if not data.startswith(b"PK\x03\x04"):
        raise ImportError_("Ese archivo no es un Excel (.xlsx) válido. Los .xls antiguos no se soportan: guárdalo como .xlsx o CSV")
    try:
        from openpyxl import load_workbook

        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        rows: list[list] = []
        for ws in wb.worksheets:  # la primera hoja que tenga datos
            for i, row in enumerate(ws.iter_rows(values_only=True)):
                if i >= MAX_ROWS + HEADER_SCAN_ROWS + 50:
                    raise ImportError_(f"La hoja tiene más de {MAX_ROWS} filas; divídela por meses")
                rows.append([_xlsx_cell(c) for c in row])
            if any(any(str(c).strip() for c in r) for r in rows):
                break
            rows = []
        wb.close()
    except ImportError_:
        raise
    except Exception as e:  # zip corrupto, hoja protegida, etc.
        raise ImportError_("No se pudo leer el Excel: el archivo parece dañado o protegido") from e
    return rows


READERS = {".csv": read_csv, ".txt": read_csv, ".tsv": read_csv, ".xlsx": read_xlsx}


def read_table(filename: str, data: bytes) -> list[list[str]]:
    ext = "." + filename.lower().rsplit(".", 1)[-1] if "." in filename else ""
    reader = READERS.get(ext)
    if not reader:
        raise ImportError_(f"Formato no soportado ({ext or 'sin extensión'}). Sube un archivo {', '.join(sorted(READERS))}")
    rows = reader(data)
    if not any(any(c for c in r) for r in rows):
        raise ImportError_("El archivo no tiene datos")
    return rows


# ---------------------------------------------------------------- encabezados y columnas
def _clean_header(h: str) -> str:
    """'Fecha de Transacción (DD/MM)' → 'fecha de transaccion dd mm'."""
    return " ".join(re.sub(r"[^a-z ]", " ", fold(str(h))).split())


def detect_mapping(header: list[str]) -> dict[str, int | None]:
    mapping: dict[str, int | None] = {f: None for f in FIELD_ORDER}
    taken: set[int] = set()
    for field_name in FIELD_ORDER:
        syns = SYNONYMS[field_name]
        for idx, raw in enumerate(header):
            h = _clean_header(raw)
            if idx in taken or not h:
                continue
            if h in syns or any(s in h for s in syns if len(s) >= 5):
                mapping[field_name] = idx
                taken.add(idx)
                break
    return mapping


def find_header(rows: list[list[str]]) -> int | None:
    """Índice de la fila de encabezados (los extractos suelen traer líneas de resumen antes), o None."""
    for i, row in enumerate(rows[:HEADER_SCAN_ROWS]):
        m = detect_mapping(row)
        if m["date"] is not None and (m["amount"] is not None or m["debit"] is not None or m["credit"] is not None):
            return i
    return None


# ---------------------------------------------------------------- montos y fechas
def parse_amount(raw: str) -> Decimal | None:
    """'$ 1.234.567,89' → 1234567.89 ; '(50.000)' → -50000 ; '50.000-' → -50000 ; '' → None."""
    s = (raw or "").strip()
    if not s:
        return None
    negative = False
    if s.startswith("(") and s.endswith(")"):
        negative, s = True, s[1:-1]
    if s.endswith("-"):
        negative, s = True, s[:-1]
    s = re.sub(r"[^\d.,+-]", "", s)  # quita $, COP, espacios
    if s.startswith("-"):
        negative, s = True, s[1:]
    s = s.lstrip("+")
    if not s or not re.search(r"\d", s):
        return None

    last_dot, last_comma = s.rfind("."), s.rfind(",")
    if last_dot != -1 and last_comma != -1:  # el separador que aparece al final es el decimal
        dec = "." if last_dot > last_comma else ","
        thou = "," if dec == "." else "."
        s = s.replace(thou, "").replace(dec, ".")
    elif last_dot != -1 or last_comma != -1:
        sep = "." if last_dot != -1 else ","
        parts = s.split(sep)
        if len(parts) > 2:  # 1.234.567 → miles
            s = "".join(parts)
        elif len(parts[1]) == 3:  # 50.000 / 1,234 → miles (en COP casi nunca hay milésimas)
            s = "".join(parts)
        else:  # 12,5 / 12.50 → decimal
            s = parts[0] + "." + parts[1]
    try:
        value = Decimal(s)
    except InvalidOperation:
        return None
    return -value if negative else value


DATE_PATTERNS = {
    "dmy": ["%d/%m/%Y", "%d-%m-%Y", "%d/%m/%y", "%d-%m-%y", "%d.%m.%Y", "%d %m %Y"],
    "mdy": ["%m/%d/%Y", "%m-%d-%Y", "%m/%d/%y", "%m-%d-%y"],
}
ISO_PATTERNS = ["%Y-%m-%d", "%Y/%m/%d", "%Y%m%d"]


def parse_date(raw, order: str = "dmy") -> date | None:
    if isinstance(raw, datetime):
        return raw.date()
    if isinstance(raw, date):
        return raw
    s = str(raw or "").strip()
    if not s:
        return None
    s = re.split(r"[ T]\d{1,2}:\d{2}", s)[0].strip()  # descarta la hora si viene
    for fmt in ISO_PATTERNS + DATE_PATTERNS.get(order, DATE_PATTERNS["dmy"]):
        try:
            d = datetime.strptime(s, fmt).date()
        except ValueError:
            continue
        if 1990 <= d.year <= clock.today().year + 1:
            return d
    return None


# ---------------------------------------------------------------- filas
@dataclass
class ParsedRow:
    line: int
    date: _dt.date | None = None
    description: str = ""
    amount: Decimal | None = None  # siempre positivo
    kind: str | None = None  # income | expense
    error: str | None = None
    import_hash: str | None = None
    extras: dict = field(default_factory=dict)


def fingerprint(account_id: int, d: date, amount: Decimal, kind: str, description: str, occurrence: int) -> str:
    base = f"{account_id}|{d.isoformat()}|{amount:.2f}|{kind}|{fold(description).strip()}|{occurrence}"
    return hashlib.sha1(base.encode()).hexdigest()


def build_rows(rows: list[list], mapping: dict[str, int | None], header_index: int | None, *,
               account_id: int, date_order: str = "dmy", invert: bool = False) -> list[ParsedRow]:
    """Convierte las filas crudas en movimientos validados (las malas quedan con `error`)."""
    if mapping["date"] is None:
        raise ImportError_("Indica cuál columna es la fecha")
    if mapping["amount"] is None and mapping["debit"] is None and mapping["credit"] is None:
        raise ImportError_("Indica la columna del monto (o las de débito y crédito)")
    start = 0 if header_index is None else header_index + 1
    out: list[ParsedRow] = []
    seen: dict[tuple, int] = {}

    def cell(row, idx):
        return row[idx] if idx is not None and idx < len(row) else ""

    for offset, row in enumerate(rows[start:]):
        if not any(str(c).strip() for c in row):
            continue
        if len(out) >= MAX_ROWS:
            raise ImportError_(f"El archivo tiene más de {MAX_ROWS} movimientos; divídelo por meses")
        pr = ParsedRow(line=start + offset + 1)
        pr.description = re.sub(r"\s+", " ", str(cell(row, mapping["description"])).strip())[:255]
        pr.date = parse_date(cell(row, mapping["date"]), date_order)
        if pr.date is None:
            pr.error = f"Fecha no válida: «{cell(row, mapping['date'])}»"
            out.append(pr)
            continue

        debit = parse_amount(str(cell(row, mapping["debit"]))) if mapping["debit"] is not None else None
        credit = parse_amount(str(cell(row, mapping["credit"]))) if mapping["credit"] is not None else None
        single = parse_amount(str(cell(row, mapping["amount"]))) if mapping["amount"] is not None else None
        signed: Decimal | None = None
        if debit or credit:
            signed = (credit or Decimal(0)) - (debit or Decimal(0))
            if debit and credit and debit == credit:
                signed = None
        elif single is not None:
            signed = single
        if not signed:
            pr.error = "Sin monto (o monto en cero)"
            out.append(pr)
            continue
        if invert:
            signed = -signed
        pr.kind = "income" if signed > 0 else "expense"
        pr.amount = abs(signed).quantize(Decimal("0.01"))
        if pr.amount > Decimal("999999999999"):
            pr.error, pr.kind, pr.amount = "Monto demasiado grande", None, None
            out.append(pr)
            continue

        key = (pr.date, pr.amount, pr.kind, fold(pr.description).strip())
        seen[key] = seen.get(key, 0) + 1  # dos cafés iguales el mismo día no son duplicado entre sí
        pr.import_hash = fingerprint(account_id, pr.date, pr.amount, pr.kind, pr.description, seen[key])
        out.append(pr)
    return out
