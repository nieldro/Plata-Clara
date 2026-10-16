"""Exportar movimientos a CSV (abre bien en Excel en español) y a Excel (.xlsx)."""
import csv
import io
from collections import defaultdict
from dataclasses import dataclass
from datetime import date
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

MAX_ROWS = 20_000
HEADERS = ["Fecha", "Cuenta", "Categoría", "Tipo", "Monto", "Nota", "Etiquetas", "Recibos"]
KIND_LABEL = {"income": "Ingreso", "expense": "Gasto", "transfer": "Transferencia", "loan": "Préstamo"}
FORMULA_STARTS = ("=", "+", "-", "@", "\t", "\r")
GOLD = "C9A96A"


@dataclass
class Row:
    date: date
    account: str
    category: str
    kind: str          # income | expense
    label: str         # Ingreso | Gasto | Transferencia | Préstamo
    amount: Decimal    # con signo: + ingreso, − gasto
    note: str
    tags: str
    receipts: int


def safe_text(value: str) -> str:
    """Un texto que empieza con = + - @ lo ejecutaría Excel como fórmula (inyección de fórmulas).
    Se antepone una comilla para que se quede como texto."""
    value = value or ""
    return "'" + value if value.startswith(FORMULA_STARTS) else value


def _decimal_es(n: Decimal) -> str:
    return f"{n:.2f}".replace(".", ",")


def to_csv(rows: list[Row]) -> bytes:
    """UTF-8 con BOM, separador «;» y coma decimal: así Excel en español lo abre sin asistentes ni símbolos raros."""
    buf = io.StringIO(newline="")
    w = csv.writer(buf, delimiter=";", lineterminator="\r\n")
    w.writerow(HEADERS)
    for r in rows:
        w.writerow([r.date.isoformat(), safe_text(r.account), safe_text(r.category), r.label, _decimal_es(r.amount),
                    safe_text(r.note), safe_text(r.tags), r.receipts or ""])
    return ("﻿" + buf.getvalue()).encode("utf-8")


def to_xlsx(rows: list[Row]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Movimientos"
    ws.append(HEADERS)
    head_fill = PatternFill("solid", fgColor=GOLD)
    for c in ws[1]:
        c.font, c.fill = Font(bold=True, color="17130A"), head_fill
    for r in rows:
        ws.append([r.date, safe_text(r.account), safe_text(r.category), r.label, float(r.amount),
                   safe_text(r.note), safe_text(r.tags), r.receipts or None])
    for row in ws.iter_rows(min_row=2):
        row[0].number_format = "yyyy-mm-dd"
        row[4].number_format = '#,##0.00;[Red]-#,##0.00'
        for cell in row:
            if isinstance(cell.value, str):
                cell.data_type = "s"  # nunca como fórmula
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for i, width in enumerate([12, 24, 22, 14, 16, 40, 22, 9], start=1):
        ws.column_dimensions[get_column_letter(i)].width = width

    # Hoja de resumen: solo cuenta ingresos y gastos reales (no transferencias ni préstamos)
    real = [r for r in rows if r.label in ("Ingreso", "Gasto")]
    income = sum((r.amount for r in real if r.kind == "income"), Decimal(0))
    expense = sum((-r.amount for r in real if r.kind == "expense"), Decimal(0))
    by_cat: dict[str, Decimal] = defaultdict(Decimal)
    for r in real:
        if r.kind == "expense":
            by_cat[r.category or "Sin categoría"] += -r.amount
    s = wb.create_sheet("Resumen")
    s.append(["Concepto", "Monto"])
    for c in s[1]:
        c.font, c.fill = Font(bold=True, color="17130A"), head_fill
    s.append(["Ingresos", float(income)])
    s.append(["Gastos", float(expense)])
    s.append(["Balance", float(income - expense)])
    s.append([])
    s.append(["Gasto por categoría", "Monto"])
    for c in s[s.max_row]:
        c.font = Font(bold=True)
    for name, total in sorted(by_cat.items(), key=lambda kv: -kv[1]):
        s.append([safe_text(name), float(total)])
    for row in s.iter_rows(min_row=2):
        row[1].number_format = "#,##0.00"
        row[0].alignment = Alignment(horizontal="left")
        if isinstance(row[0].value, str):
            row[0].data_type = "s"
    s.column_dimensions["A"].width = 30
    s.column_dimensions["B"].width = 18

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
